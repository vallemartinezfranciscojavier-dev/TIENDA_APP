"""
Servicio de exportación de reportes (Fase 13).

Exporta únicamente el reporte activo (ventas, productos, ganancias,
inventario, compras, fiado, caja o movimientos) a Excel (.xlsx) y PDF.

Consume la misma estructura que genera la vista:

    {
        "reporte": clave,
        "titulo": str,
        "periodo_label": str,
        "fecha_generacion": datetime,
        "resumen": [(etiqueta, valor, tipo), ...],
        "tabla": {encabezados, filas, columnas_moneda, columnas_cantidad, columnas_enteros},
        "tablas_extra": [ {misma forma que tabla}, ... ],
    }

Dependencias: openpyxl (Excel) y reportlab (PDF).
"""

from datetime import datetime

from utils.formatters import formatear_fecha, formatear_moneda


def _etiqueta(datos: dict) -> str:
    return datos.get("periodo_label", "")


def _titulo_documento(datos: dict) -> str:
    """Nombre del documento: tienda + reporte (sin duplicar "Reporte")."""
    from config.settings import APP_NAME
    titulo = (datos.get("titulo") or "").strip()
    if not titulo:
        titulo = "Reporte"
    elif not titulo.lower().startswith("reporte"):
        titulo = f"Reporte {titulo}"
    return f"{APP_NAME} — {titulo}"


def _a_float(valor):
    """Convertir números (incl. Decimal) a float; dejar lo demás igual."""
    from decimal import Decimal

    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, (int, float)):
        return float(valor)
    return valor


def _secciones(datos: dict):
    """Todas las secciones de tabla del reporte (principal + extra)."""
    secciones = []
    if datos.get("tabla"):
        secciones.append(datos["tabla"])
    secciones.extend(datos.get("tablas_extra", []) or [])
    return secciones


def _valor_celda(valor, c, seccion):
    """Formatear una celda para PDF según el tipo de columna."""
    moneda = seccion.get("columnas_moneda", []) or []
    cantidad = seccion.get("columnas_cantidad", []) or []
    enteros = seccion.get("columnas_enteros", []) or []
    if c in moneda:
        return formatear_moneda(valor)
    if c in enteros:
        try:
            return str(int(valor or 0))
        except (TypeError, ValueError):
            return str(valor)
    if c in cantidad:
        try:
            from decimal import Decimal
            s = format(Decimal(str(valor or 0)), "f")
            if "." in s:
                s = s.rstrip("0").rstrip(".")
            return s if s else "0"
        except Exception:
            return str(valor)
    return str(valor if valor is not None else "")


# ======================================================================
# Excel
# ======================================================================

def _estilo_excel():
    from openpyxl.styles import Font, PatternFill

    return {
        "titulo": Font(bold=True, size=14, color="2C3E50"),
        "encabezado": Font(bold=True, color="FFFFFF"),
        "relleno_encabezado": PatternFill("solid", fgColor="2C3E50"),
        "seccion": Font(bold=True, size=11, color="2C3E50"),
        "currency_fmt": "#,##0.00",
    }


def _agregar_tabla_excel(ws, titulo, encabezados, filas, est, columnas_moneda):
    """
    Escribir una tabla con título en la hoja: título, encabezado y filas.
    Las columnas indicadas en ``columnas_moneda`` reciben formato de moneda.
    """
    fila_titulo = ws.max_row + 1
    if titulo:
        ws.cell(row=fila_titulo, column=1, value=titulo).font = est["seccion"]
        fila_titulo += 1

    ws.append(encabezados)
    fila_encabezado = ws.max_row
    for celda in ws[fila_encabezado]:
        if celda.value is None:
            continue
        celda.font = est["encabezado"]
        celda.fill = est["relleno_encabezado"]

    for fila in filas:
        ws.append([_a_float(v) for v in fila])
    for col in (columnas_moneda or []):
        for f in range(fila_encabezado + 1, ws.max_row + 1):
            celda = ws.cell(row=f, column=col + 1)  # openpyxl es 1-based
            if isinstance(celda.value, (int, float)):
                celda.number_format = est["currency_fmt"]


def _ajustar_anchos(ws, anchos):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho


def exportar_reporte_excel(ruta, datos: dict):
    """
    Exportar el reporte activo a un archivo Excel.
    Hoja 1 resumen; hoja(s) siguientes con los datos del reporte.
    """
    from openpyxl import Workbook

    est = _estilo_excel()
    wb = Workbook()

    # --------------------------------------------------------------
    # Hoja Resumen
    # --------------------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 22

    ws["A1"] = f"{_titulo_documento(datos)} — {_etiqueta(datos)}"
    ws["A1"].font = est["titulo"]
    ws["A2"] = f"Generado: {formatear_fecha(datos.get('fecha_generacion', datetime.now()), '%d/%m/%Y %H:%M')}"

    fila = 4
    for etiqueta, valor, tipo in datos.get("resumen", []):
        ws.cell(row=fila, column=1, value=str(etiqueta)).font = est["seccion"]
        valor_num = _a_float(valor)
        celda = ws.cell(row=fila, column=2, value=valor_num)
        if tipo == "moneda" and isinstance(valor_num, (int, float)):
            celda.number_format = est["currency_fmt"]
        if tipo == "porcentaje" and isinstance(valor_num, (int, float)):
            celda.number_format = '0.0"%"'
        fila += 1

    # --------------------------------------------------------------
    # Hojas de datos
    # --------------------------------------------------------------
    secciones = _secciones(datos)
    for idx, seccion in enumerate(secciones, start=2):
        titulo_hoja = (
            (seccion.get("titulo") or f"Reporte")
            if idx == 2
            else (seccion.get("titulo") or f"Sección {idx - 1}")
        )
        hoja = wb.create_sheet(titulo_hoja[:31])
        _agregar_tabla_excel(
            hoja,
            "",
            seccion.get("encabezados", []),
            seccion.get("filas", []),
            est,
            seccion.get("columnas_moneda", []),
        )
        _ajustar_anchos(hoja, [max(10, min(32, len(str(cab)) + 12)) for cab in seccion.get("encabezados", [])])

    wb.save(ruta)


# ======================================================================
# PDF
# ======================================================================

def exportar_reporte_pdf(ruta, datos: dict):
    """
    Exportar el reporte activo a un PDF con tablas legibles.

    Incluye: título, periodo, fecha de generación, resumen superior y los
    datos del reporte activo (no mezcla los 8 reportes).
    """
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    estilos = getSampleStyleSheet()
    titulo_style = ParagraphStyle(
        "Titulo", parent=estilos["Title"], fontSize=16, textColor=colors.HexColor("#2C3E50")
    )
    seccion_style = ParagraphStyle(
        "Seccion", parent=estilos["Heading2"], fontSize=12, textColor=colors.HexColor("#2980B9"),
        spaceBefore=14, spaceAfter=6,
    )
    normal = estilos["Normal"]
    resumen_style = ParagraphStyle(
        "ResumenCelda", parent=estilos["Normal"], fontSize=9, leading=11,
    )

    CAB = colors.HexColor("#2C3E50")

    def _tabla(seccion):
        encabezados = seccion.get("encabezados", [])
        filas = [encabezados]
        for fila in seccion.get("filas", []):
            filas.append([
                _valor_celda(valor, c, seccion) for c, valor in enumerate(fila)
            ])
        t = Table(filas, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), CAB),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 8.5),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F3F5")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D5DB")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    doc = SimpleDocTemplate(
        str(ruta),
        pagesize=landscape(letter),
        leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28,
        title=_titulo_documento(datos),
    )

    historia = []
    historia.append(Paragraph(_titulo_documento(datos), titulo_style))
    historia.append(Paragraph(
        f"<b>Periodo:</b> {_etiqueta(datos)} &nbsp;&nbsp; "
        f"<b>Generado:</b> {formatear_fecha(datos.get('fecha_generacion', datetime.now()), '%d/%m/%Y %H:%M')}",
        normal,
    ))
    historia.append(Spacer(1, 10))

    # --- Resumen (tarjetas) ---
    resumen = datos.get("resumen", [])
    if resumen:
        historia.append(Paragraph("Resumen", seccion_style))
        # Cada celda del resumen lleva su propia etiqueta en línea.
        # Se usa Paragraph para que reportlab renderice el markup <b>.
        celdas_resumen = []
        for etiqueta, valor, tipo in resumen:
            celdas_resumen.append(
                Paragraph(f"<b>{etiqueta}:</b> {_valor_tarjeta(etiqueta, valor, tipo)}", resumen_style)
            )
        pares = []
        for i in range(0, len(celdas_resumen), 2):
            par = [celdas_resumen[i]]
            if i + 1 < len(celdas_resumen):
                par.append(celdas_resumen[i + 1])
            else:
                par.append("")
            pares.append(par)
        t_resumen = Table(pares)
        t_resumen.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D5DB")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D0D5DB")),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        historia.append(t_resumen)
        historia.append(Spacer(1, 8))

    # --- Tablas ---
    for idx, seccion in enumerate(_secciones(datos)):
        titulo = seccion.get("titulo") or "Detalle"
        historia.append(Paragraph(titulo, seccion_style))
        filas_datos = seccion.get("filas", [])
        if filas_datos:
            historia.append(_tabla(seccion))
        else:
            historia.append(Paragraph("Sin datos para mostrar.", normal))
        historia.append(Spacer(1, 6))

    doc.build(historia)


def _valor_tarjeta(etiqueta, valor, tipo) -> str:
    """Valor formateado de una tarjeta de resumen para el PDF."""
    if tipo == "moneda":
        return formatear_moneda(valor)
    if tipo == "porcentaje":
        from utils.formatters import formatear_porcentaje
        return formatear_porcentaje(valor)
    if tipo == "cantidad":
        return _valor_celda(valor, -1, {
            "columnas_moneda": [], "columnas_cantidad": [-1], "columnas_enteros": [],
        })
    if tipo == "numero":
        try:
            return str(int(valor or 0))
        except (TypeError, ValueError):
            return str(valor)
    return str(valor)